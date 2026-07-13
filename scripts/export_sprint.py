#!/usr/bin/env python3
"""
Export Alibaba Yunxiao (center-edition) sprint work items to Excel.

Uses the shared yunxiao_client library for API access. This script only
handles Excel generation and CLI arguments.

Targeting a sprint (in priority order):
  1. --sprint-id         direct sprint identifier
  2. config sprint_id    direct sprint identifier from config
  3. --sprint-keyword / config sprint_keyword
                         filter sprints whose name contains the keyword; if any
                         of the matched sprints has a trailing YYMMDD date, pick
                         the one closest to today; otherwise prefer a DOING
                         sprint, then the most recently modified, then the first.

Usage:
    python scripts/export_sprint.py --kind all
    python scripts/export_sprint.py --sprint-keyword "招聘" --kind req
    python scripts/export_sprint.py --sprint-id 06175757e51bfd3f2aa47f3e60 --kind task
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any

# Add lib/ to path
SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(SKILL_DIR / "lib"))

from yunxiao_client import (
    DEFAULT_DOMAIN,
    DEFAULT_SPRINT_STATUSES,
    YunxiaoClient,
    build_config_from_sources,
    column_value,
    custom_value,
    discover_headers,
    first_real,
    load_config_file,
    resolve_sprint,
    standard_value,
    STANDARD_FIELD_BY_HEADER,
    filter_excluded_statuses,
)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

STANDARD_REQUIREMENT_HEADERS = ["编号", "标题", "描述", "负责人", "状态", "参与者", "工作项类型"]
STANDARD_TASK_HEADERS = ["编号", "标题", "描述", "负责人", "状态", "工作项类型"]

DEFAULT_EXCEL = {
    "requirement_sheet_title": "当前迭代需求",
    "task_sheet_title": "当前迭代任务",
    "requirement_headers": [],
    "task_headers": [],
}

CONFIG_DIR = SKILL_DIR / "config"
DEFAULT_CONFIG_PATH = CONFIG_DIR / "yunxiao_config.json"
EXAMPLE_CONFIG_PATH = CONFIG_DIR / "yunxiao_config.example.json"


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------

def build_rows(items: list[dict[str, Any]], headers: list[str]) -> list[list[str]]:
    return [[column_value(item, header) for header in headers] for item in items]


def _format_number(value: float) -> str:
    if value == int(value):
        return str(int(value))
    return f"{value:.2f}"


def resolve_summary_sum_fields(available_headers: list[str]) -> list[str]:
    if "预计工时" in available_headers:
        return ["预计工时"]
    return []


def build_summary(
    items: list[dict[str, Any]],
    group_header: str,
    sum_headers: list[str],
) -> tuple[list[str], list[list[str]]]:
    groups: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for item in items:
        name = column_value(item, group_header) or "(未分配)"
        if name not in groups:
            groups[name] = {"count": 0, "sums": {f: 0.0 for f in sum_headers}}
            order.append(name)
        groups[name]["count"] += 1
        for field in sum_headers:
            raw = column_value(item, field)
            try:
                groups[name]["sums"][field] += float(str(raw).strip() or 0)
            except (ValueError, TypeError):
                pass

    headers = [group_header, "任务数", *sum_headers]
    rows: list[list[str]] = []
    total_count = 0
    total_sums = {f: 0.0 for f in sum_headers}
    for name in order:
        data = groups[name]
        rows.append([name, str(data["count"])] + [_format_number(data["sums"][f]) for f in sum_headers])
        total_count += data["count"]
        for f in sum_headers:
            total_sums[f] += data["sums"][f]
    rows.append(["总计", str(total_count)] + [_format_number(total_sums[f]) for f in sum_headers])
    return headers, rows


def render_summary_markdown(headers: list[str], rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(headers) + " |"]
    lines.append("|" + "|".join(["---"] * len(headers)) + "|")
    for row in rows:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[str]], metadata: dict[str, str]) -> None:
    if len(workbook.sheetnames) == 1 and workbook.active.max_row == 1 and workbook.active.max_column == 1:
        sheet = workbook.active
    else:
        sheet = workbook.create_sheet()
    sheet.title = title[:31]

    sheet.append([title])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A1"].alignment = Alignment(horizontal="center")

    for key, value in metadata.items():
        sheet.append([key, value])

    header_row = sheet.max_row + 2
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(row)

    widths = [16, 40, 70, 18, 16, 24, 24]
    for index, width in enumerate(widths[: len(headers)], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, sheet.max_row)}"


def save_workbook(workbook: Workbook, path: Path) -> None:
    workbook.save(path)


def write_excel(path: Path, title: str, headers: list[str], rows: list[list[str]], metadata: dict[str, str]) -> None:
    workbook = Workbook()
    write_sheet(workbook, title, headers, rows, metadata)
    save_workbook(workbook, path)


def dump_samples(label: str, items: list[dict[str, Any]], count: int) -> None:
    print(f"=== {label} samples (showing {min(count, len(items))} of {len(items)}) ===")
    if not items:
        print("(no items in this sprint)")
        return
    for index, item in enumerate(items[:count], start=1):
        custom = item.get("customFieldValues") or item.get("customFields") or []
        serial = item.get("serialNumber") or item.get("identifier") or f"#{index}"
        print(f"\n--- {label} sample {index}: {serial} ---")
        print(json.dumps(custom, ensure_ascii=False, indent=2))


# ---------------------------------------------------------------------------
# Config and CLI
# ---------------------------------------------------------------------------

def build_config(args: argparse.Namespace) -> dict[str, Any]:
    config_file = Path(args.config) if args.config else DEFAULT_CONFIG_PATH
    file_config = load_config_file(config_file)

    cli_args = {
        "domain": args.domain,
        "token": args.token,
        "organization_id": args.organization_id,
        "project_id": args.project_id,
        "sprint_id": args.sprint_id,
        "sprint_keyword": args.sprint_keyword,
        "sprint_statuses": args.sprint_statuses,
        "corp_id": args.corp_id,
    }
    config = build_config_from_sources(file_config, cli_args)

    if not config["sprint_id"] and not config["sprint_keyword"]:
        raise RuntimeError(
            "Specify either sprint_id or sprint_keyword (CLI flag, env var, or config file)."
        )

    excel_user = file_config.get("excel") if isinstance(file_config.get("excel"), dict) else {}
    config["excel"] = {**DEFAULT_EXCEL, **excel_user}
    return config


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Yunxiao center-edition sprint work items to Excel.")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG_PATH), help=f"Config JSON path. Default: {DEFAULT_CONFIG_PATH}")
    parser.add_argument("--domain", help=f"Yunxiao domain. Default: {DEFAULT_DOMAIN}")
    parser.add_argument("--token", help="Yunxiao personal access token. Prefer YUNXIAO_TOKEN.")
    parser.add_argument("--organization-id", help="Yunxiao organizationId.")
    parser.add_argument("--project-id", help="Yunxiao projectId.")
    parser.add_argument("--corp-id", help="Optional corpId for output metadata.")
    parser.add_argument("--sprint-id", help="Direct sprint id to export.")
    parser.add_argument("--sprint-keyword", help="Keyword to filter sprint names (contains match).")
    parser.add_argument("--sprint-statuses", default=DEFAULT_SPRINT_STATUSES, help=f"Comma-separated sprint statuses. Default: {DEFAULT_SPRINT_STATUSES}")
    parser.add_argument("--today", help="Override today for date-based selection, format YYYY-MM-DD.")
    parser.add_argument("--kind", choices=["req", "task", "all"], default="all")
    parser.add_argument("--output-dir", default="exports")
    parser.add_argument("--dump-sample", type=int, default=0, metavar="N", help="Print N sample workitems' customFieldValues as JSON and exit.")
    parser.add_argument("--show-config-path", action="store_true", help="Print the absolute config file path and exit.")
    return parser.parse_args()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> int:
    args = parse_args()

    if args.show_config_path:
        config_file = Path(args.config) if args.config else DEFAULT_CONFIG_PATH
        print(f"Config file: {config_file}")
        print(f"Exists:      {'yes' if config_file.exists() else 'no (copy from example below)'}")
        if EXAMPLE_CONFIG_PATH.exists():
            print(f"Template:    {EXAMPLE_CONFIG_PATH}")
            print(f"Bootstrap:   cp {EXAMPLE_CONFIG_PATH} {config_file}")
        return 0

    config = build_config(args)
    client = YunxiaoClient(config)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    today = dt.date.today()
    if args.today:
        today = dt.datetime.strptime(args.today, "%Y-%m-%d").date()

    if config["sprint_id"]:
        sprint = {"id": config["sprint_id"], "name": args.sprint_id or config["sprint_id"]}
    else:
        sprints = client.list_sprints(config["project_id"], config["sprint_statuses"])
        sprint = resolve_sprint(sprints, config["sprint_keyword"], today)

    sprint_id = str(sprint.get("id") or sprint.get("identifier") or "")
    sprint_name = str(sprint.get("name") or sprint.get("displayName") or sprint_id)
    if not sprint_id:
        raise RuntimeError(f"Cannot determine sprint id from sprint: {sprint}")

    if args.dump_sample > 0:
        if args.kind in ("req", "all"):
            req_items = client.fetch_all_workitems(config["project_id"], sprint_id, "Req")
            dump_samples("Requirements", req_items, args.dump_sample)
        if args.kind in ("task", "all"):
            task_items = client.fetch_all_workitems(config["project_id"], sprint_id, "Task")
            dump_samples("Tasks", task_items, args.dump_sample)
        print("\n(tip) Copy the fieldName values above into excel.task_headers / requirement_headers.")
        return 0

    metadata = {
        "导出时间": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "组织ID": config["organization_id"],
        "CorpID": config["corp_id"],
        "项目ID": config["project_id"],
        "迭代": sprint_name,
        "迭代ID": sprint_id,
    }
    metadata = {key: value for key, value in metadata.items() if value}

    excel = config["excel"]
    generated: list[Path] = []
    safe_sprint_name = re.sub(r"[^\w.-]+", "_", sprint_name)

    if args.kind in ("req", "all"):
        req_items = client.fetch_all_workitems(config["project_id"], sprint_id, "Req")
        req_items = filter_excluded_statuses(req_items)
        print(f"Requirements: fetching descriptions for {len(req_items)} items...")
        client.enrich_with_descriptions(req_items)
        req_path = output_dir / f"{safe_sprint_name}_requirements.xlsx"
        configured_req_headers = list(excel["requirement_headers"])
        req_headers = configured_req_headers or discover_headers(req_items, STANDARD_REQUIREMENT_HEADERS)
        if not configured_req_headers:
            print(f"Requirements: auto-discovered headers ({len(req_headers)} cols)")
        write_excel(
            req_path,
            excel["requirement_sheet_title"],
            req_headers,
            build_rows(req_items, req_headers),
            metadata,
        )
        generated.append(req_path)
        print(f"Requirements: {len(req_items)} -> {req_path.resolve()}")

    task_summary_md: str | None = None
    if args.kind in ("task", "all"):
        task_items = client.fetch_all_workitems(config["project_id"], sprint_id, "Task")
        task_items = filter_excluded_statuses(task_items)
        print(f"Tasks: fetching descriptions for {len(task_items)} items...")
        client.enrich_with_descriptions(task_items)
        task_path = output_dir / f"{safe_sprint_name}_tasks.xlsx"
        configured_task_headers = list(excel["task_headers"])
        task_headers = configured_task_headers or discover_headers(task_items, STANDARD_TASK_HEADERS)
        if not configured_task_headers:
            print(f"Tasks: auto-discovered headers ({len(task_headers)} cols)")

        workbook = Workbook()
        write_sheet(workbook, excel["task_sheet_title"], task_headers, build_rows(task_items, task_headers), metadata)

        sum_fields = resolve_summary_sum_fields(task_headers)
        if sum_fields:
            group_by = "负责人" if "负责人" in task_headers else None
        if sum_fields and group_by:
            summary_headers, summary_rows = build_summary(task_items, group_by, sum_fields)
            write_sheet(workbook, "工时汇总", summary_headers, summary_rows, metadata)
            task_summary_md = render_summary_markdown(summary_headers, summary_rows)
            print(f"Tasks: summary sheet added (group by {group_by}, sum {sum_fields})")
        else:
            print("Tasks: summary skipped (need '负责人' column and '预计工时' field in task headers)")

        save_workbook(workbook, task_path)
        generated.append(task_path)
        print(f"Tasks: {len(task_items)} -> {task_path.resolve()}")

    print("\nGenerated files (absolute paths):")
    for path in generated:
        print(f"  {path.resolve()}")

    if task_summary_md:
        print("\n=== 工时汇总 (also written to sheet '工时汇总') ===")
        print(task_summary_md)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
