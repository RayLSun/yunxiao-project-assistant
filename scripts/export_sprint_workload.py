#!/usr/bin/env python3
"""
Export sprint workload (work hours) statistics for team R&D members.

Fetches all requirements in a sprint, then finds their child tasks,
filters by team assignees, and aggregates 预计工时 per person.
Outputs an Excel file with per-person summary + task detail.

Usage:
    python scripts/export_sprint_workload.py --sprint-name "Sprint260728"
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path
from typing import Any

# Add lib/ to path
SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(SKILL_DIR / "lib"))

from yunxiao_client import (
    DEFAULT_DOMAIN,
    DEFAULT_SPRINT_STATUSES,
    YunxiaoClient,
    build_config_from_sources,
    first_real,
    load_config_file,
    load_team,
    resolve_sprint,
    object_name,
    filter_excluded_statuses,
)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CONFIG_DIR = SKILL_DIR / "config"
DEFAULT_CONFIG_PATH = CONFIG_DIR / "yunxiao_config.json"
EXAMPLE_CONFIG_PATH = CONFIG_DIR / "yunxiao_config.example.json"
OUTPUT_DIR = SKILL_DIR / "doc" / "workload"


# ---------------------------------------------------------------------------
# Config and CLI
# ---------------------------------------------------------------------------

def build_config(args: argparse.Namespace) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    config_file = Path(args.config) if args.config else DEFAULT_CONFIG_PATH
    file_config = load_config_file(config_file)
    cli_args = {
        "domain": args.domain,
        "token": args.token,
        "organization_id": args.organization_id,
        "project_id": args.project_id,
    }
    config = build_config_from_sources(file_config, cli_args)
    team_members = load_team(file_config)
    return config, team_members


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export sprint workload statistics for team R&D members.")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG_PATH))
    parser.add_argument("--domain", help=f"Yunxiao domain. Default: {DEFAULT_DOMAIN}")
    parser.add_argument("--token", help="Yunxiao personal access token.")
    parser.add_argument("--organization-id", help="Yunxiao organizationId.")
    parser.add_argument("--project-id", help="Yunxiao projectId.")
    parser.add_argument("--sprint-name", required=True, help="Sprint name (exact or keyword match).")
    parser.add_argument("--sprint-statuses", default=DEFAULT_SPRINT_STATUSES, help=f"Default: {DEFAULT_SPRINT_STATUSES}")
    parser.add_argument("--today", help="Override today, format YYYY-MM-DD.")
    parser.add_argument("--show-config-path", action="store_true")
    return parser.parse_args()


# ---------------------------------------------------------------------------
# Workload calculation
# ---------------------------------------------------------------------------

def extract_hours(item: dict[str, Any]) -> float:
    """Extract 预计工时 from a workitem's customFieldValues."""
    for cf in item.get("customFieldValues") or []:
        if cf.get("fieldName") == "预计工时":
            vals = cf.get("values") or []
            if vals:
                v = vals[0]
                if isinstance(v, dict):
                    h = v.get("displayValue") or v.get("identifier")
                else:
                    h = v
                try:
                    return float(str(h).strip())
                except (ValueError, TypeError):
                    pass
    return 0.0


def extract_assignee(item: dict[str, Any]) -> tuple[str, str]:
    """Return (user_id, display_name) from assignedTo."""
    assigned = item.get("assignedTo")
    if isinstance(assigned, dict):
        uid = str(assigned.get("id") or assigned.get("identifier") or "")
        name = str(assigned.get("displayName") or assigned.get("name") or "")
        return uid, name
    return "", ""


def resolve_sprint_by_name(client: YunxiaoClient, project_id: str, sprint_name: str, statuses: str) -> dict[str, Any]:
    """Find a sprint by exact name, then keyword match."""
    sprints = client.list_sprints(project_id, statuses)
    needle = sprint_name.lower()
    exact = [s for s in sprints if str(s.get("name") or s.get("displayName") or "").lower() == needle]
    if exact:
        return exact[0]
    matches = [s for s in sprints if needle in str(s.get("name") or s.get("displayName") or "").lower()]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        names = [str(s.get("name") or s.get("displayName") or "") for s in matches]
        raise RuntimeError(f"Multiple sprints match '{sprint_name}': {', '.join(names)}")
    raise RuntimeError(f"No sprint found matching '{sprint_name}'")


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

def write_excel(
    output_path: Path,
    sprint_name: str,
    sprint_id: str,
    config: dict[str, Any],
    team_members: list[dict[str, Any]],
    summary_rows: list[list[str]],
    detail_rows: list[list[str]],
) -> None:
    workbook = Workbook()

    # --- Summary sheet ---
    summary_sheet = workbook.active
    summary_sheet.title = "工时汇总"

    summary_headers = ["姓名", "角色", "TL", "任务数", "预计工时", "要求工时", "差异", "状态"]
    metadata = {
        "导出时间": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "组织ID": config["organization_id"],
        "项目ID": config["project_id"],
        "迭代": sprint_name,
        "迭代ID": sprint_id,
        "统计范围": "需求关联的子任务，按负责人汇总",
    }

    # Title
    summary_sheet.append([f"{sprint_name} - 研发工时统计"])
    summary_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(summary_headers))
    summary_sheet["A1"].font = Font(bold=True, size=14)
    summary_sheet["A1"].alignment = Alignment(horizontal="center")

    for key, value in metadata.items():
        summary_sheet.append([key, value])

    header_row = summary_sheet.max_row + 2
    for col, header in enumerate(summary_headers, start=1):
        cell = summary_sheet.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in summary_rows:
        summary_sheet.append(row)

    # Color the 状态 column
    status_col = len(summary_headers)
    for row_idx in range(header_row + 1, summary_sheet.max_row + 1):
        cell = summary_sheet.cell(row=row_idx, column=status_col)
        val = str(cell.value or "")
        if "不足" in val:
            cell.fill = PatternFill("solid", fgColor="FFC7CE")
            cell.font = Font(color="9C0006")
        elif "超出" in val:
            cell.fill = PatternFill("solid", fgColor="C6EFCE")
            cell.font = Font(color="006100")
        elif "达标" in val:
            cell.fill = PatternFill("solid", fgColor="C6EFCE")
            cell.font = Font(color="006100")

    # Widths
    widths = [10, 12, 6, 8, 10, 10, 8, 10]
    for i, w in enumerate(widths, start=1):
        summary_sheet.column_dimensions[get_column_letter(i)].width = w

    summary_sheet.freeze_panes = summary_sheet.cell(row=header_row + 1, column=1)

    # --- Detail sheet ---
    detail_sheet = workbook.create_sheet("任务明细")
    detail_headers = ["编号", "标题", "负责人", "角色", "预计工时", "状态", "父需求编号", "父需求标题"]

    detail_sheet.append([f"{sprint_name} - 任务明细"])
    detail_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(detail_headers))
    detail_sheet["A1"].font = Font(bold=True, size=14)
    detail_sheet["A1"].alignment = Alignment(horizontal="center")

    header_row2 = detail_sheet.max_row + 2
    for col, header in enumerate(detail_headers, start=1):
        cell = detail_sheet.cell(row=header_row2, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in detail_rows:
        detail_sheet.append(row)

    for row in detail_sheet.iter_rows(min_row=header_row2 + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    detail_widths = [16, 50, 10, 12, 10, 12, 16, 50]
    for i, w in enumerate(detail_widths, start=1):
        detail_sheet.column_dimensions[get_column_letter(i)].width = w

    detail_sheet.freeze_panes = detail_sheet.cell(row=header_row2 + 1, column=1)
    detail_sheet.auto_filter.ref = f"A{header_row2}:{get_column_letter(len(detail_headers))}{max(header_row2, detail_sheet.max_row)}"

    workbook.save(output_path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> int:
    args = parse_args()

    if args.show_config_path:
        config_file = Path(args.config) if args.config else DEFAULT_CONFIG_PATH
        print(f"Config file: {config_file}")
        print(f"Exists:      {'yes' if config_file.exists() else 'no'}")
        return 0

    config, team_members = build_config(args)
    client = YunxiaoClient(config)

    # Build team lookup: yunxiao_id -> member dict
    team_by_id: dict[str, dict[str, Any]] = {}
    for m in team_members:
        if m["yunxiao_id"] and not m["exempt"]:
            team_by_id[m["yunxiao_id"]] = m

    if not team_by_id:
        print("ERROR: No non-exempt team members with yunxiao_id found.", file=sys.stderr)
        return 1

    print(f"Team R&D members (non-exempt): {len(team_by_id)}", file=sys.stderr)

    # Resolve sprint
    sprint = resolve_sprint_by_name(client, config["project_id"], args.sprint_name, args.sprint_statuses)
    sprint_id = str(sprint.get("id") or sprint.get("identifier") or "")
    sprint_name = str(sprint.get("name") or sprint.get("displayName") or sprint_id)
    if not sprint_id:
        print(f"ERROR: Cannot determine sprint id from sprint: {sprint}", file=sys.stderr)
        return 1
    print(f"Sprint: {sprint_name} ({sprint_id})", file=sys.stderr)

    # Fetch requirements in this sprint (to get parent IDs)
    print("Fetching requirements...", file=sys.stderr)
    reqs = client.fetch_all_workitems(config["project_id"], sprint_id, "Req")
    reqs = filter_excluded_statuses(reqs)
    req_ids = set(str(r.get("id") or "") for r in reqs)
    req_by_id: dict[str, dict[str, Any]] = {str(r.get("id") or ""): r for r in reqs}
    print(f"  {len(reqs)} requirements found.", file=sys.stderr)

    # Fetch ALL tasks in project (not sprint-filtered) so we don't miss
    # tasks whose sprint differs from their parent requirement's sprint
    print("Fetching tasks...", file=sys.stderr)
    tasks = client.fetch_all_workitems(config["project_id"], "", "Task")
    tasks = filter_excluded_statuses(tasks)
    print(f"  {len(tasks)} tasks found.", file=sys.stderr)

    # Filter: task's parent must be a requirement in this sprint
    # AND assignee must be a team R&D member
    detail_rows: list[list[str]] = []
    hours_by_person: dict[str, dict[str, Any]] = {}

    for t in tasks:
        pid = str(t.get("parentId") or "")
        if pid not in req_ids:
            continue

        uid, name = extract_assignee(t)
        if uid not in team_by_id:
            continue

        member = team_by_id[uid]
        hours = extract_hours(t)

        # Detail row
        parent_req = req_by_id.get(pid, {})
        detail_rows.append([
            str(t.get("serialNumber") or ""),
            str(t.get("subject") or ""),
            name,
            member["role"],
            f"{hours:.0f}" if hours == int(hours) else f"{hours:.1f}",
            object_name(t.get("status")),
            str(parent_req.get("serialNumber") or ""),
            str(parent_req.get("subject") or ""),
        ])

        # Aggregate
        if uid not in hours_by_person:
            hours_by_person[uid] = {
                "name": name,
                "role": member["role"],
                "is_tl": member["is_tl"],
                "required_hours": member["required_hours"],
                "hours": 0.0,
                "task_count": 0,
            }
        hours_by_person[uid]["hours"] += hours
        hours_by_person[uid]["task_count"] += 1

    # Build summary rows - follow config team.members order (not alphabetical)
    summary_rows: list[list[str]] = []
    member_order = [m["yunxiao_id"] for m in team_members if m.get("yunxiao_id")]
    for uid in member_order:
        if uid not in hours_by_person:
            continue
        d = hours_by_person[uid]
        diff = d["hours"] - d["required_hours"]
        if abs(diff) < 0.5:
            status = "达标"
        elif diff < 0:
            status = f"不足 {abs(diff):.0f}h"
        else:
            status = f"超出 {diff:.0f}h"
        summary_rows.append([
            d["name"],
            d["role"],
            "Yes" if d["is_tl"] else "",
            str(d["task_count"]),
            f"{d['hours']:.0f}" if d["hours"] == int(d["hours"]) else f"{d['hours']:.1f}",
            str(d["required_hours"]),
            f"{diff:+.0f}",
            status,
        ])

    # Also include team members who have 0 hours (no tasks in this sprint's requirements)
    # Follow config team.members order for zero-hour members too (non-exempt only)
    zero_hour_added = {r[0] for r in summary_rows}  # already added names
    for m in team_members:
        uid = m.get("yunxiao_id", "")
        if uid and uid not in hours_by_person and not m.get("exempt", False) and m["name"] not in zero_hour_added:
            summary_rows.append([
                m["name"],
                m["role"],
                "Yes" if m["is_tl"] else "",
                "0",
                "0",
                str(m["required_hours"]),
                f"{-m['required_hours']:+.0f}",
                f"不足 {m['required_hours']}h",
            ])

    # Sort detail rows by assignee name then serial number
    detail_rows.sort(key=lambda r: (r[2], r[0]))

    # Create output
    safe_sprint_name = re.sub(r"[^\w.-]+", "_", sprint_name)
    sprint_output_dir = OUTPUT_DIR / safe_sprint_name
    sprint_output_dir.mkdir(parents=True, exist_ok=True)
    output_path = sprint_output_dir / f"{safe_sprint_name}_workload.xlsx"

    write_excel(output_path, sprint_name, sprint_id, config, team_members, summary_rows, detail_rows)

    # Print summary to stdout
    print(f"\n=== {sprint_name} 工时统计 ===\n")
    print(f"{'姓名':<8} {'角色':<12} {'TL':<4} {'任务数':<6} {'预计':<6} {'要求':<6} {'差异':<6} 状态")
    print("-" * 75)
    for row in summary_rows:
        name, role, tl, tc, hours, req, diff, status = row
        print(f"{name:<8} {role:<12} {tl:<4} {tc:<6} {hours:<6} {req:<6} {diff:<6} {status}")
    print(f"\nExported to {output_path.resolve()}")

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
